from __future__ import annotations

import asyncio
import hmac
import json
import io
import logging
import os
import re
import threading
import uuid
import urllib.error
import urllib.request
import zipfile
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from docx import Document
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile, Query
from pypdf import PdfReader
from pydantic import BaseModel, Field


from ai_lead_desk import lead_store

PROJECT_DIR = Path(__file__).resolve().parents[2]
DOCUMENTS_DIR = PROJECT_DIR / "data" / "institute_documents"
UPLOADS_DIR = DOCUMENTS_DIR / "uploads"

load_dotenv(PROJECT_DIR / ".env")

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

logger = logging.getLogger("ai_lead_desk.backend")

app = FastAPI(
    title="AI Lead Desk Backend",
    description="Counselor document management and RAG indexing API.",
    version="1.0.0",
)

_rag_service = None
_rag_lock = threading.Lock()


def require_counselor_key(
    x_api_key: str = Header(default="", alias="X-API-Key"),
) -> None:
    """Protect counselor endpoints using a locally configured API key."""
    expected = os.getenv("COUNSELOR_API_KEY", "").strip()

    if not expected:
        raise HTTPException(
            status_code=503,
            detail="Counselor API authentication is not configured.",
        )

    if not hmac.compare_digest(x_api_key, expected):
        raise HTTPException(
            status_code=401,
            detail="Invalid or missing counselor API key.",
        )


def _safe_filename(original_name: str | None) -> tuple[str, str]:
    """Return a safe display filename and its validated extension."""
    raw_name = (original_name or "").replace("\\", "/").rsplit("/", 1)[-1]
    raw_name = raw_name.strip()

    if not raw_name or raw_name in {".", ".."}:
        raise HTTPException(status_code=400, detail="A filename is required.")

    extension = Path(raw_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Only PDF and DOCX documents are supported.",
        )

    stem = Path(raw_name).stem
    safe_stem = re.sub(r"[^A-Za-z0-9 _.-]", "_", stem)
    safe_stem = safe_stem.strip(" ._")[:100] or "document"

    return f"{safe_stem}{extension}", extension


def _validate_document(filename: str, content: bytes) -> None:
    """Verify file structure and ensure the file contains extractable text."""
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    extension = Path(filename).suffix.lower()

    try:
        if extension == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise HTTPException(
                    status_code=422,
                    detail="The file does not have a valid PDF signature.",
                )

            reader = PdfReader(io.BytesIO(content), strict=False)

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=422,
                    detail="Password-protected PDFs are not supported.",
                )

            extracted_text = "\n".join(
                page.extract_text() or "" for page in reader.pages
            )

            if not extracted_text.strip():
                raise HTTPException(
                    status_code=422,
                    detail=(
                        "No selectable text was found in this PDF. "
                        "Scanned documents requiring OCR are not supported yet."
                    ),
                )

        elif extension == ".docx":
            stream = io.BytesIO(content)

            if not zipfile.is_zipfile(stream):
                raise HTTPException(
                    status_code=422,
                    detail="The uploaded file is not a valid DOCX document.",
                )

            stream.seek(0)
            with zipfile.ZipFile(stream) as archive:
                names = set(archive.namelist())
                if not {
                    "[Content_Types].xml",
                    "word/document.xml",
                }.issubset(names):
                    raise HTTPException(
                        status_code=422,
                        detail="The DOCX document structure is invalid.",
                    )

                entries = archive.infolist()
                uncompressed_size = sum(item.file_size for item in entries)

                if len(entries) > 2500 or uncompressed_size > 50 * 1024 * 1024:
                    raise HTTPException(
                        status_code=413,
                        detail="The DOCX archive exceeds safe processing limits.",
                    )

            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )

            if not "\n".join(parts).strip():
                raise HTTPException(
                    status_code=422,
                    detail="No extractable text was found in this DOCX.",
                )

    except HTTPException:
        raise
    except Exception:
        logger.warning(
            "Document validation failed for extension %s.",
            extension,
        )
        raise HTTPException(
            status_code=422,
            detail="The document is invalid or could not be read.",
        ) from None


def _rebuild_rag_index():
    """Build a fresh index and publish it only after construction succeeds."""
    global _rag_service

    from ai_lead_desk.rag_service import DocumentRAGService

    with _rag_lock:
        new_service = DocumentRAGService(project_dir=PROJECT_DIR)
        _rag_service = new_service

    return new_service


def _list_documents() -> list[dict]:
    """List existing institute documents without exposing absolute paths."""
    if not DOCUMENTS_DIR.is_dir():
        return []

    with _rag_lock:
        indexed_sources = (
            {item.get("source") for item in _rag_service.document_chunks}
            if _rag_service is not None
            else None
        )

    documents = []

    for path in DOCUMENTS_DIR.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {
            ".pdf", ".docx", ".txt"
        }:
            continue

        try:
            relative_path = path.relative_to(DOCUMENTS_DIR).as_posix()
            upload_id = None

            try:
                upload_relative = path.relative_to(UPLOADS_DIR)
                if len(upload_relative.parts) >= 2:
                    upload_id = upload_relative.parts[0]
            except ValueError:
                pass

            if indexed_sources is None:
                index_status = "not_checked"
            elif path.name in indexed_sources:
                index_status = "indexed"
            else:
                index_status = "pending"

            stat = path.stat()
            documents.append({
                "document_id": upload_id,
                "filename": path.name,
                "relative_path": relative_path,
                "extension": path.suffix.lower(),
                "size_bytes": stat.st_size,
                "modified_at": datetime.fromtimestamp(
                    stat.st_mtime
                ).astimezone().isoformat(timespec="seconds"),
                "index_status": index_status,
            })
        except OSError:
            logger.warning("Could not inspect a stored document.")

    documents.sort(key=lambda item: item["filename"].lower())
    return documents


@app.get("/health")
def health():
    """Public liveness check; reveals no credentials or document content."""
    return {"status": "ok", "service": "ai-lead-desk-backend"}


@app.post("/api/documents/upload", status_code=201)
async def upload_document(
    file: UploadFile = File(...),
    _: None = Depends(require_counselor_key),
):
    """Validate and store an institute PDF or DOCX, then attempt re-indexing."""
    filename, _extension = _safe_filename(file.filename)

    try:
        content = await file.read(MAX_UPLOAD_BYTES + 1)
    finally:
        await file.close()

    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail="The maximum upload size is 15 MB.",
        )

    _validate_document(filename, content)

    document_id = uuid.uuid4().hex
    document_folder = UPLOADS_DIR / document_id
    document_path = document_folder / filename
    temporary_path = document_folder / ".upload.partial"

    try:
        document_folder.mkdir(parents=True, exist_ok=False)
        temporary_path.write_bytes(content)
        temporary_path.replace(document_path)
    except OSError:
        import shutil
        shutil.rmtree(document_folder, ignore_errors=True)
        logger.error("Unable to store an uploaded document.")
        raise HTTPException(
            status_code=500,
            detail="The document could not be stored.",
        ) from None

    index_status = "indexed"
    index_message = "Document stored and indexed."

    try:
        service = await asyncio.to_thread(_rebuild_rag_index)
        if not any(
            chunk.get("source") == filename
            for chunk in service.document_chunks
        ):
            index_status = "pending"
            index_message = (
                "Document was stored, but no index chunks were produced."
            )
    except Exception as exc:
        logger.warning(
            "Document stored but re-indexing failed: %s",
            type(exc).__name__,
        )
        index_status = "pending"
        index_message = (
            "Document was stored, but indexing could not complete. "
            "Retry using the reindex endpoint."
        )

    return {
        "document_id": document_id,
        "filename": filename,
        "size_bytes": len(content),
        "index_status": index_status,
        "message": index_message,
    }


@app.get("/api/documents")
def list_documents(
    _: None = Depends(require_counselor_key),
):
    """List currently stored institute documents."""
    documents = _list_documents()
    return {"count": len(documents), "documents": documents}


@app.post("/api/documents/reindex")
async def reindex_documents(
    _: None = Depends(require_counselor_key),
):
    """Rebuild the complete RAG index from the stored institute documents."""
    try:
        service = await asyncio.to_thread(_rebuild_rag_index)
    except Exception as exc:
        logger.warning(
            "RAG index rebuild failed: %s",
            type(exc).__name__,
        )
        raise HTTPException(
            status_code=503,
            detail=(
                "The document index could not be rebuilt. "
                "Check the backend model and embedding services."
            ),
        ) from None

    return {
        "status": "indexed",
        "documents": len({
            item["source"] for item in service.document_chunks
        }),
        "chunks": len(service.document_chunks),
    }

# =========================================================
# COUNSELOR CONVERSATION INBOX
# =========================================================

class CounselorReplyRequest(BaseModel):
    content: str = Field(min_length=1, max_length=4000)
    counselor: str = Field(default="counselor", min_length=1, max_length=100)


def _require_connected_platform(platform: str) -> str:
    """Limit conversation actions to currently integrated platforms."""
    normalized = str(platform).strip().lower()

    if normalized != "telegram":
        raise HTTPException(
            status_code=501,
            detail="This messaging platform is not integrated yet.",
        )

    return normalized


def _get_existing_lead(platform: str, platform_user_id: str) -> dict:
    """Fetch an existing lead without creating an empty conversation."""
    platform = str(platform).strip().lower()
    platform_user_id = str(platform_user_id).strip()

    if not platform_user_id:
        raise HTTPException(status_code=400, detail="A user ID is required.")

    with lead_store._connect() as connection:
        row = connection.execute(
            """
            SELECT *
            FROM leads
            WHERE lower(platform) = ? AND platform_user_id = ?
            """,
            (platform, platform_user_id),
        ).fetchone()

    if row is None:
        raise HTTPException(status_code=404, detail="Conversation not found.")

    return dict(row)


def _record_counselor_event(
    lead_id: int,
    event_type: str,
    actor: str = "counselor",
    details: str | None = None,
) -> None:
    """Write a minimal audit event without duplicating message content."""
    with lead_store._connect() as connection:
        connection.execute(
            """
            INSERT INTO counselor_events (
                lead_id, event_type, actor, details, created_at
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                lead_id,
                event_type,
                str(actor)[:100],
                details,
                lead_store._now(),
            ),
        )


def _send_telegram_message(chat_id: str, text: str) -> int:
    """Send a counselor reply without logging or returning the bot token."""
    token = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()

    if not token:
        raise HTTPException(
            status_code=503,
            detail="Telegram sending is not configured on the backend.",
        )

    payload = json.dumps({
        "chat_id": chat_id,
        "text": text,
    }).encode("utf-8")

    request = urllib.request.Request(
        f"https://api.telegram.org/bot{token}/sendMessage",
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            result = json.loads(response.read().decode("utf-8"))

    except urllib.error.HTTPError as exc:
        # Do not log the exception URL because Telegram URLs contain the token.
        logger.warning(
            "Telegram rejected a counselor reply (HTTP %s).",
            exc.code,
        )
        raise HTTPException(
            status_code=502,
            detail="Telegram rejected the counselor reply.",
        ) from None

    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        logger.warning(
            "Telegram connection failed during counselor reply (%s).",
            type(exc).__name__,
        )
        raise HTTPException(
            status_code=502,
            detail="Could not reach Telegram to send the reply.",
        ) from None

    except (ValueError, UnicodeDecodeError):
        raise HTTPException(
            status_code=502,
            detail="Telegram returned an unreadable response.",
        ) from None

    if not result.get("ok") or not isinstance(result.get("result"), dict):
        logger.warning("Telegram returned an unsuccessful send result.")
        raise HTTPException(
            status_code=502,
            detail="Telegram could not confirm delivery of the reply.",
        )

    return int(result["result"]["message_id"])


@app.get("/api/conversations")
def list_conversations(
    platform: str = Query(default="telegram", min_length=1, max_length=32),
    limit: int = Query(default=100, ge=1, le=500),
    _: None = Depends(require_counselor_key),
):
    """Return inbox summaries for the selected messaging platform."""
    normalized = _require_connected_platform(platform)
    conversations = lead_store.list_conversations(
        limit=limit,
        platform=normalized,
    )
    return {
        "count": len(conversations),
        "conversations": conversations,
    }


@app.get("/api/conversations/{platform}/{platform_user_id}/messages")
def get_conversation_messages(
    platform: str,
    platform_user_id: str,
    limit: int = Query(default=200, ge=1, le=500),
    _: None = Depends(require_counselor_key),
):
    """Return chronological persisted messages for the selected student."""
    normalized = _require_connected_platform(platform)
    lead = _get_existing_lead(normalized, platform_user_id)

    with lead_store._connect() as connection:
        rows = connection.execute(
            """
            SELECT id, direction, content, created_at, sender_type
            FROM lead_messages
            WHERE lead_id = ?
            ORDER BY id DESC
            LIMIT ?
            """,
            (lead["id"], limit),
        ).fetchall()

    messages = [dict(row) for row in reversed(rows)]

    return {
        "conversation": {
            "lead_id": lead["id"],
            "platform": lead["platform"],
            "platform_user_id": lead["platform_user_id"],
            "name": lead.get("name"),
            "handling_mode": lead.get("handling_mode") or "ai",
            "assigned_counselor": lead.get("assigned_counselor"),
            "chat_ready": bool(lead.get("chat_id")),
        },
        "count": len(messages),
        "messages": messages,
    }


@app.post("/api/conversations/{platform}/{platform_user_id}/read")
def mark_conversation_as_read(
    platform: str,
    platform_user_id: str,
    _: None = Depends(require_counselor_key),
):
    """Mark the messages currently in a conversation as seen by a counselor."""
    normalized = _require_connected_platform(platform)
    _get_existing_lead(normalized, platform_user_id)

    latest_id = lead_store.mark_conversation_read(
        normalized,
        platform_user_id,
    )

    return {
        "status": "ok",
        "last_read_message_id": latest_id,
    }


@app.post("/api/conversations/{platform}/{platform_user_id}/takeover")
def takeover_conversation(
    platform: str,
    platform_user_id: str,
    _: None = Depends(require_counselor_key),
):
    """Give a conversation to a human counselor."""
    normalized = _require_connected_platform(platform)
    lead = _get_existing_lead(normalized, platform_user_id)

    updated = lead_store.set_handling_mode(
        normalized,
        platform_user_id,
        "human",
        actor="counselor",
    )

    return {
        "status": "ok",
        "platform": normalized,
        "platform_user_id": str(platform_user_id),
        "handling_mode": updated.get("handling_mode", "human"),
        "chat_ready": bool(updated.get("chat_id")),
        "message": (
            "Human takeover enabled."
            if updated.get("chat_id")
            else
            "Human takeover enabled. A Telegram chat destination has not "
            "been recorded yet; the student needs to message the bot first."
        ),
    }


@app.post("/api/conversations/{platform}/{platform_user_id}/resume")
def resume_ai_conversation(
    platform: str,
    platform_user_id: str,
    _: None = Depends(require_counselor_key),
):
    """Return a conversation to AI-controlled handling."""
    normalized = _require_connected_platform(platform)
    _get_existing_lead(normalized, platform_user_id)

    updated = lead_store.set_handling_mode(
        normalized,
        platform_user_id,
        "ai",
        actor="counselor",
    )

    return {
        "status": "ok",
        "platform": normalized,
        "platform_user_id": str(platform_user_id),
        "handling_mode": updated.get("handling_mode", "ai"),
        "message": "AI handling resumed.",
    }


@app.post("/api/conversations/{platform}/{platform_user_id}/reply")
async def send_counselor_reply(
    platform: str,
    platform_user_id: str,
    payload: CounselorReplyRequest,
    _: None = Depends(require_counselor_key),
):
    """Send a counselor reply to the user's original Telegram chat."""
    normalized = _require_connected_platform(platform)
    lead = _get_existing_lead(normalized, platform_user_id)

    if (lead.get("handling_mode") or "ai") != "human":
        raise HTTPException(
            status_code=409,
            detail="Take over this conversation before replying as a counselor.",
        )

    chat_id = lead.get("chat_id")

    if not chat_id:
        raise HTTPException(
            status_code=409,
            detail=(
                "No Telegram chat destination is stored for this lead. "
                "Ask the student to message the bot again first."
            ),
        )

    content = payload.content.strip()
    actor = payload.counselor.strip()

    if not content:
        raise HTTPException(
            status_code=422,
            detail="Reply text cannot be blank.",
        )

    if not actor:
        raise HTTPException(
            status_code=422,
            detail="Counselor name cannot be blank.",
        )

    if len(content) > 4000:
        raise HTTPException(
            status_code=422,
            detail="Replies must be 4000 characters or fewer.",
        )

    telegram_message_id = await asyncio.to_thread(
        _send_telegram_message,
        str(chat_id),
        content,
    )

    lead_store.add_message(
        normalized,
        platform_user_id,
        "outbound",
        content,
        sender_type="counselor",
    )

    _record_counselor_event(
        lead["id"],
        "reply_sent",
        actor=actor,
        details=f"telegram_message_id={telegram_message_id}",
    )

    return {
        "status": "sent",
        "platform": normalized,
        "platform_user_id": str(platform_user_id),
        "telegram_message_id": telegram_message_id,
        "message": "Counselor reply sent to Telegram.",
    }


@app.get("/api/conversations/{platform}/{platform_user_id}/events")
def get_conversation_events(
    platform: str,
    platform_user_id: str,
    limit: int = Query(default=100, ge=1, le=500),
    _: None = Depends(require_counselor_key),
):
    """Return the audit trail of handoffs and counselor actions."""
    normalized = _require_connected_platform(platform)
    _get_existing_lead(normalized, platform_user_id)

    events = lead_store.get_counselor_events(
        normalized,
        platform_user_id,
        limit=limit,
    )

    return {
        "count": len(events),
        "events": events,
    }

